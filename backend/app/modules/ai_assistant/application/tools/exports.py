"""Filtered .xlsx exports of the partner's recruiting lists (assistant tools).

Four read-only tools — ``export_jobs``, ``export_interviews``, ``export_offers``,
``export_events`` — that turn an org-scoped, service-layer-RBAC'd list read into
a REAL Excel file stored as a :class:`ChatExportFile` (24 h expiry) and returned
to the frontend as the existing ``{"kind": "download", ...}`` render artifact.
Raw bytes/storage never reach the client; only the owner-scoped
``/ai/chat/exports/{id}`` URL does.

Shared here: ONE xlsx writer (same styling as the applicant export), ONE store
helper, ONE metadata-only analytics fact per generated file (kind + row count —
never row content). Every data read happens in the owning module's application
layer (``job_export_service`` / ``recruiting_export_service`` /
``event_export_service``), which enforce the export grants:

- ``export_jobs``       → ``jobs:export``
- ``export_interviews`` → ``applications:export`` + the board's ``interviews:read``
- ``export_offers``     → ``applications:export`` + the board's ``offers:create``
  (salary column rides the same grant that unlocks comp on the offer board)
- ``export_events``     → ``events:export`` (aggregate attendee counts, no PII)
"""

from __future__ import annotations

import io
import uuid as _uuid
from datetime import UTC, datetime, time, timedelta

from sqlalchemy.ext.asyncio import AsyncSession

from app.shared.exceptions import AuthRequiredError, PermissionDeniedError
from app.shared.permissions import Principal

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_EXPORT_TTL_HOURS = 24

# --------------------------------------------------------------------------- #
# Column catalogs (ordered; localized VN-first headers)                        #
# --------------------------------------------------------------------------- #

_JOB_COLUMNS: dict[str, dict[str, str]] = {
    "title": {"vi": "Tin tuyển dụng", "en": "Job title"},
    "status": {"vi": "Trạng thái", "en": "Status"},
    "applicants": {"vi": "Số ứng viên", "en": "Applicants"},
    "unreviewed": {"vi": "Chưa xem", "en": "Unreviewed"},
    "deadline": {"vi": "Hạn ứng tuyển", "en": "Deadline"},
    "created_at": {"vi": "Ngày tạo", "en": "Created at"},
    "owner": {"vi": "Người phụ trách", "en": "Owner"},
}

_INTERVIEW_COLUMNS: dict[str, dict[str, str]] = {
    "candidate": {"vi": "Ứng viên", "en": "Candidate"},
    "job": {"vi": "Tin tuyển dụng", "en": "Job"},
    "stage": {"vi": "Vòng", "en": "Stage"},
    "scheduled_at": {"vi": "Thời gian phỏng vấn", "en": "Scheduled at"},
    "mode": {"vi": "Hình thức", "en": "Mode"},
    "interviewers": {"vi": "Người phỏng vấn", "en": "Interviewers"},
    "status": {"vi": "Trạng thái", "en": "Status"},
}

_OFFER_COLUMNS: dict[str, dict[str, str]] = {
    "candidate": {"vi": "Ứng viên", "en": "Candidate"},
    "job": {"vi": "Tin tuyển dụng", "en": "Job"},
    "position": {"vi": "Vị trí", "en": "Position"},
    "status": {"vi": "Trạng thái", "en": "Status"},
    "salary": {"vi": "Lương đề nghị", "en": "Offered salary"},
    "sent_at": {"vi": "Ngày gửi", "en": "Sent at"},
    "decided_at": {"vi": "Ngày phản hồi", "en": "Decided at"},
}

_EVENT_COLUMNS: dict[str, dict[str, str]] = {
    "title": {"vi": "Sự kiện", "en": "Event"},
    "event_type": {"vi": "Loại sự kiện", "en": "Event type"},
    "format": {"vi": "Hình thức", "en": "Format"},
    "status": {"vi": "Trạng thái", "en": "Status"},
    "starts_at": {"vi": "Bắt đầu", "en": "Starts at"},
    "ends_at": {"vi": "Kết thúc", "en": "Ends at"},
    "capacity": {"vi": "Sức chứa", "en": "Capacity"},
    "registered": {"vi": "Đã đăng ký", "en": "Registered"},
    "waitlisted": {"vi": "Danh sách chờ", "en": "Waitlisted"},
    "attended": {"vi": "Đã tham dự", "en": "Attended"},
    "created_at": {"vi": "Ngày tạo", "en": "Created at"},
}


# --------------------------------------------------------------------------- #
# Shared helpers (single xlsx writer, store, artifact, analytics fact)         #
# --------------------------------------------------------------------------- #


def _parse_columns(raw: object) -> list[str] | None:
    if isinstance(raw, str):
        raw = [c.strip() for c in raw.split(",") if c.strip()]
    if isinstance(raw, list):
        return [str(c) for c in raw]
    return None


def _resolve_columns(
    requested: list[str] | None, catalog: dict[str, dict[str, str]]
) -> list[str]:
    resolved = [c for c in (requested or []) if c in catalog]
    return resolved or list(catalog)


def _parse_date(raw: object, *, end_of_day: bool = False) -> datetime | None:
    """Parse a user-supplied YYYY-MM-DD / ISO string to an aware UTC datetime.

    ``None``/blank → ``None``; malformed → ``ValueError`` (mapped to the
    deterministic ``invalid_date`` tool error by the callers).
    """

    clean = str(raw or "").strip()
    if not clean:
        return None
    dt = datetime.fromisoformat(clean)  # raises ValueError when malformed
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=UTC)
    if end_of_day and len(clean) == 10:  # date-only upper bound is inclusive
        dt = datetime.combine(dt.date(), time.max, tzinfo=dt.tzinfo)
    return dt


def _xlsx_bytes(
    columns: list[str],
    catalog: dict[str, dict[str, str]],
    rows: list[dict],
    *,
    sheet_title: str,
    locale: str,
) -> bytes:
    """Render rows to a styled workbook — the ONE styling implementation."""

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Export"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2D5FA6")
    for idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=catalog[col].get(locale, col))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 22
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, col in enumerate(columns, start=1):
            ws.cell(row=r_idx, column=c_idx, value=row.get(col, ""))
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _store_and_render(
    session: AsyncSession,
    principal: Principal,
    *,
    export_kind: str,
    filename: str,
    content: bytes,
    row_count: int,
    columns: list[str],
    extra: dict | None = None,
) -> dict:
    """Persist the file, record the metadata-only fact, return the tool result."""

    from app.modules.ai_assistant.domain.models import ChatExportFile
    from app.modules.analytics.application import ingestion_service as analytics

    export = ChatExportFile(
        id=_uuid.uuid4(),
        user_id=principal.user_id,
        org_id=principal.org_id,
        filename=filename,
        mime=_XLSX_MIME,
        content=content,
        row_count=row_count,
        expires_at=datetime.now(UTC) + timedelta(hours=_EXPORT_TTL_HOURS),
    )
    session.add(export)

    # Metadata-only product fact: which export kind, how many rows — never row
    # content, identities, or file bytes.
    await analytics.record_event_safe(
        session,
        event_type="ai.chat_export.generated",
        aggregate_type="chat_export",
        aggregate_id=export.id,
        actor_id=principal.user_id,
        actor_type="partner",
        properties={"export": export_kind, "row_count": row_count, "columns": columns},
    )

    result = {
        "ok": True,
        "row_count": row_count,
        "columns": columns,
        "filename": filename,
        # FE-only render artifact (stripped before the model sees the result).
        "render": {
            "kind": "download",
            "download_path": f"/ai/chat/exports/{export.id}",
            "filename": filename,
            "row_count": row_count,
            "format": "xlsx",
        },
    }
    if extra:
        result.update(extra)
    return result


def _parse_uuid(raw: object) -> _uuid.UUID | None:
    if not raw:
        return None
    try:
        return _uuid.UUID(str(raw).strip())
    except ValueError:
        return None


def _stamp() -> str:
    return datetime.now(UTC).strftime("%Y%m%d")


# --------------------------------------------------------------------------- #
# Tools                                                                        #
# --------------------------------------------------------------------------- #


async def export_jobs(session: AsyncSession, principal: Principal, args: dict) -> dict:
    """Own-org job list → .xlsx (status + created-date filters, column subset)."""
    from app.modules.opportunities.application import job_export_service

    if not principal.is_authenticated or principal.org_id is None:
        return {"ok": False, "error": "partner_only"}
    try:
        created_from = _parse_date(args.get("created_from"))
        created_to = _parse_date(args.get("created_to"), end_of_day=True)
    except ValueError:
        return {"ok": False, "error": "invalid_date"}
    locale = str(args.get("locale") or "vi")

    try:
        rows = await job_export_service.export_job_rows(
            session,
            principal=principal,
            status=args.get("status"),
            created_from=created_from,
            created_to=created_to,
            locale=locale,
        )
    except PermissionDeniedError:
        return {"ok": False, "error": "permission_denied"}
    except AuthRequiredError:
        return {"ok": False, "error": "partner_only"}
    except Exception:
        return {"ok": False, "error": "tool_failed"}

    columns = _resolve_columns(_parse_columns(args.get("columns")), _JOB_COLUMNS)
    content = _xlsx_bytes(columns, _JOB_COLUMNS, rows, sheet_title="Jobs", locale=locale)
    return await _store_and_render(
        session,
        principal,
        export_kind="jobs",
        filename=f"tin-tuyen-dung-{_stamp()}.xlsx",
        content=content,
        row_count=len(rows),
        columns=columns,
    )


async def export_interviews(session: AsyncSession, principal: Principal, args: dict) -> dict:
    """Org-wide interview board → .xlsx (scope/job/date filters, no meeting links)."""
    from app.modules.recruitment.application import recruiting_export_service

    if not principal.is_authenticated or principal.org_id is None:
        return {"ok": False, "error": "partner_only"}
    try:
        scheduled_from = _parse_date(args.get("scheduled_from"))
        scheduled_to = _parse_date(args.get("scheduled_to"), end_of_day=True)
    except ValueError:
        return {"ok": False, "error": "invalid_date"}
    locale = str(args.get("locale") or "vi")

    try:
        rows = await recruiting_export_service.export_interview_rows(
            session,
            principal=principal,
            scope=str(args.get("scope") or "all"),
            job_id=_parse_uuid(args.get("job_id")),
            scheduled_from=scheduled_from,
            scheduled_to=scheduled_to,
            locale=locale,
        )
    except PermissionDeniedError:
        return {"ok": False, "error": "permission_denied"}
    except AuthRequiredError:
        return {"ok": False, "error": "partner_only"}
    except Exception:
        return {"ok": False, "error": "tool_failed"}

    columns = _resolve_columns(_parse_columns(args.get("columns")), _INTERVIEW_COLUMNS)
    content = _xlsx_bytes(
        columns, _INTERVIEW_COLUMNS, rows, sheet_title="Interviews", locale=locale
    )
    return await _store_and_render(
        session,
        principal,
        export_kind="interviews",
        filename=f"lich-phong-van-{_stamp()}.xlsx",
        content=content,
        row_count=len(rows),
        columns=columns,
    )


async def export_offers(session: AsyncSession, principal: Principal, args: dict) -> dict:
    """Org-wide offer board → .xlsx (status filter; salary rides the comp grant)."""
    from app.modules.recruitment.application import recruiting_export_service

    if not principal.is_authenticated or principal.org_id is None:
        return {"ok": False, "error": "partner_only"}
    locale = str(args.get("locale") or "vi")

    try:
        rows, salary_included = await recruiting_export_service.export_offer_rows(
            session,
            principal=principal,
            status=(str(args.get("status")).strip() or None) if args.get("status") else None,
            job_id=_parse_uuid(args.get("job_id")),
            locale=locale,
        )
    except PermissionDeniedError:
        return {"ok": False, "error": "permission_denied"}
    except AuthRequiredError:
        return {"ok": False, "error": "partner_only"}
    except Exception:
        return {"ok": False, "error": "tool_failed"}

    columns = _resolve_columns(_parse_columns(args.get("columns")), _OFFER_COLUMNS)
    if not salary_included:
        columns = [c for c in columns if c != "salary"]
    content = _xlsx_bytes(columns, _OFFER_COLUMNS, rows, sheet_title="Offers", locale=locale)
    return await _store_and_render(
        session,
        principal,
        export_kind="offers",
        filename=f"thu-moi-nhan-viec-{_stamp()}.xlsx",
        content=content,
        row_count=len(rows),
        columns=columns,
        extra={"salary_included": salary_included},
    )


async def export_events(session: AsyncSession, principal: Principal, args: dict) -> dict:
    """Own-org events → .xlsx (aggregate attendee counts only — never attendee PII)."""
    from app.modules.opportunities.application import event_export_service

    if not principal.is_authenticated or principal.org_id is None:
        return {"ok": False, "error": "partner_only"}
    locale = str(args.get("locale") or "vi")

    try:
        rows = await event_export_service.export_event_rows(
            session,
            principal=principal,
            status=args.get("status"),
            locale=locale,
        )
    except PermissionDeniedError:
        return {"ok": False, "error": "permission_denied"}
    except AuthRequiredError:
        return {"ok": False, "error": "partner_only"}
    except Exception:
        return {"ok": False, "error": "tool_failed"}

    columns = _resolve_columns(_parse_columns(args.get("columns")), _EVENT_COLUMNS)
    content = _xlsx_bytes(columns, _EVENT_COLUMNS, rows, sheet_title="Events", locale=locale)
    return await _store_and_render(
        session,
        principal,
        export_kind="events",
        filename=f"su-kien-{_stamp()}.xlsx",
        content=content,
        row_count=len(rows),
        columns=columns,
    )
