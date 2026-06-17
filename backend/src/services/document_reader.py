"""Text extraction helpers for uploaded JD/CV files."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path


def extract_text_from_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".docx", ".docs"}:
        return _extract_docx(path)
    raise ValueError("Supported uploaded files: .txt, .pdf, .docx")


def extract_text_from_bytes(file_name: str, content: bytes) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".txt":
        return content.decode("utf-8")
    if suffix == ".pdf":
        return _extract_pdf_bytes(content)
    if suffix in {".docx", ".docs"}:
        return _extract_docx_bytes(content)
    raise ValueError("Supported uploaded files: .txt, .pdf, .docx")


def _extract_pdf(path: Path) -> str:
    text = _extract_pdf_with_pypdf(path)
    if _has_enough_text(text):
        return text
    text = _extract_pdf_with_pdfplumber(path)
    return _ensure_pdf_text(text)


def _extract_pdf_bytes(content: bytes) -> str:
    data = BytesIO(content)
    text = _extract_pdf_with_pypdf(data)
    if _has_enough_text(text):
        return text
    data.seek(0)
    text = _extract_pdf_with_pdfplumber(data)
    return _ensure_pdf_text(text)


def _extract_pdf_with_pypdf(source: Path | BytesIO) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Install pypdf to read PDF files.") from exc

    reader = PdfReader(source)
    return "\n".join(page.extract_text() or "" for page in reader.pages).strip()


def _extract_pdf_with_pdfplumber(source: Path | BytesIO) -> str:
    try:
        import pdfplumber
    except ImportError:
        return ""

    with pdfplumber.open(source) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages).strip()


def _has_enough_text(text: str) -> bool:
    compact = " ".join(text.split())
    return len(compact) >= 12 and len(compact.split()) >= 3


def _ensure_pdf_text(text: str) -> str:
    if _has_enough_text(text):
        return text
    raise ValueError(
        "Không đọc được đủ chữ từ PDF. File có thể là CV scan/ảnh hoặc dùng biểu đồ kỹ năng dạng hình; cần OCR hoặc model vision để đọc chính xác."
    )


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Install python-docx to read DOCX files.") from exc

    doc = Document(str(path))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs).strip()


def _extract_docx_bytes(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Install python-docx to read DOCX files.") from exc

    doc = Document(BytesIO(content))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs).strip()
