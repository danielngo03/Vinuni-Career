"""Locally-generated CV ingestion fixtures (``docs/CV_INGESTION_EXTRACTION_SPEC.md`` §8).

Fixtures are GENERATED at runtime (deterministic, no committed binaries) so the
eval set runs anywhere the lightweight deps (fpdf2 / python-docx / PIL) are
available. Each builder returns ``bytes``.

Covered cases: text PDF (en), Vietnamese text CV, two-column PDF, image/scanned
(mocked OCR), sparse/canvas PDF with an embedded image, DOCX, blank PDF, non-CV
PDF, corrupt PDF, and mixed vi/en. Password-protected and duplicate are exercised
via a mocked extractor error / re-upload in the tests.
"""

from __future__ import annotations

import io


def _pdf_bytes(builder) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    builder(pdf)
    return bytes(pdf.output())


def _line(pdf, text: str) -> None:
    """Write one text line and return the cursor to the left margin."""

    from fpdf.enums import XPos, YPos

    if text:
        pdf.multi_cell(0, 7, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    else:
        pdf.ln(4)


# --------------------------------------------------------------------------- #
# Text-selectable PDFs                                                          #
# --------------------------------------------------------------------------- #


def text_pdf_en() -> bytes:
    def build(pdf) -> None:
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        lines = [
            "Jane Engineer",
            "Email: jane.engineer@example.com | Phone: +84 912 000 111",
            "",
            "Summary",
            "Backend engineering intern focused on APIs and data.",
            "",
            "Education",
            "VinUniversity - BSc Computer Science, 2022 - 2026.",
            "",
            "Experience",
            "Software Intern, Example Tech (2024 - 2025).",
            "Built REST APIs with FastAPI and PostgreSQL.",
            "",
            "Skills",
            "Python, FastAPI, SQLAlchemy, PostgreSQL.",
        ]
        for line in lines:
            _line(pdf, line)

    return _pdf_bytes(build)


def two_column_pdf() -> bytes:
    def build(pdf) -> None:
        pdf.add_page()
        pdf.set_font("helvetica", size=11)
        # Left column.
        pdf.set_xy(10, 20)
        pdf.multi_cell(
            85,
            6,
            "Alex Twocolumn\nEmail: alex.two@example.com\n\nSkills\n"
            "Python\nFastAPI\nSQL\n\nLanguages\nEnglish\nVietnamese",
        )
        # Right column.
        pdf.set_xy(105, 20)
        pdf.multi_cell(
            85,
            6,
            "Experience\nSoftware Intern at Example (2024-2025)\n"
            "Built internal APIs.\n\nEducation\nBSc Computer Science, 2026.",
        )

    return _pdf_bytes(build)


def sparse_canvas_pdf() -> bytes:
    """Vector/canvas-heavy PDF with an embedded image and almost no native text.

    pdfplumber reports embedded images + sub-threshold text, so the cascade routes
    this to the OCR stage (scanned-document path).
    """

    def build(pdf) -> None:
        pdf.add_page()
        # Vector content only.
        pdf.set_line_width(0.5)
        for i in range(0, 180, 12):
            pdf.line(15, 20 + i, 195, 20 + i)
        # Embed a small raster image so page.images is non-empty.
        pdf.image(io.BytesIO(_png_bytes()), x=20, y=20, w=40, h=40)

    return _pdf_bytes(build)


def blank_pdf() -> bytes:
    def build(pdf) -> None:
        pdf.add_page()  # empty page, no text, no images

    return _pdf_bytes(build)


def not_cv_pdf() -> bytes:
    def build(pdf) -> None:
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        _line(
            pdf,
            "Meeting minutes for the neighborhood committee. Today we discussed "
            "the plan to organize a party and decorate the venue for the upcoming "
            "festival next week. No further business was raised and the meeting "
            "was adjourned at five in the afternoon.",
        )

    return _pdf_bytes(build)


def corrupt_pdf() -> bytes:
    return b"%PDF-1.4 this is not really a valid pdf body at all"


# --------------------------------------------------------------------------- #
# DOCX                                                                          #
# --------------------------------------------------------------------------- #


def docx_cv() -> bytes:
    import docx

    document = docx.Document()
    for line in [
        "Minh Nguyen",
        "Email: minh.nguyen@example.com | Phone: 0901234567",
        "Summary",
        "Aspiring data engineer with internship experience.",
        "Education",
        "VinUniversity - BSc Computer Science, 2026.",
        "Experience",
        "Data Intern, Example Corp (2024-2025). Built ETL pipelines in Python.",
        "Skills",
        "Python, SQL, Airflow, dbt.",
    ]:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def vietnamese_docx_cv() -> bytes:
    import docx

    document = docx.Document()
    for line in [
        "Nguyễn Văn An",
        "Email: an.nguyen@example.com | Điện thoại: 0987654321",
        "Mục tiêu",
        "Thực tập sinh kỹ thuật phần mềm, tập trung vào phát triển API.",
        "Học vấn",
        "Trường Đại học VinUni - Cử nhân Khoa học Máy tính, 2026.",
        "Kinh nghiệm",
        "Thực tập sinh phần mềm tại Example (2024-2025). Xây dựng API bằng FastAPI.",
        "Kỹ năng",
        "Python, FastAPI, SQL, PostgreSQL.",
    ]:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Text fixtures (utf-8) for language detection                                 #
# --------------------------------------------------------------------------- #


def vietnamese_cv_txt() -> bytes:
    return (
        "Nguyễn Thị Bình\n"
        "Email: binh.nguyen@example.com | Điện thoại: 0912345678\n\n"
        "Mục tiêu\nThực tập sinh phân tích dữ liệu.\n\n"
        "Học vấn\nTrường Đại học VinUni - Khoa học Máy tính, 2026.\n\n"
        "Kinh nghiệm\nThực tập tại công ty Example, xây dựng báo cáo dữ liệu.\n\n"
        "Kỹ năng\nPython, SQL, Excel.\n"
    ).encode()


def mixed_language_cv_txt() -> bytes:
    return (
        "Tran Mixed Candidate\n"
        "Email: mixed.candidate@example.com | Phone: +84 900 111 222\n\n"
        "Mục tiêu / Objective\n"
        "Software engineering intern. Thực tập sinh kỹ thuật phần mềm.\n\n"
        "Kinh nghiệm / Experience\n"
        "Built REST APIs with FastAPI. Xây dựng API với FastAPI và PostgreSQL.\n\n"
        "Kỹ năng / Skills\nPython, FastAPI, SQL, kỹ năng làm việc nhóm.\n"
    ).encode()


# --------------------------------------------------------------------------- #
# Raster image fixtures                                                         #
# --------------------------------------------------------------------------- #


def _png_bytes(text: str = "CV") -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (240, 120), "white")
    draw = ImageDraw.Draw(image)
    draw.text((10, 50), text, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def scanned_image_cv() -> bytes:
    """A raster image upload (no native text). Tests pair this with a mocked OCR."""

    return _png_bytes("Scanned CV")


__all__ = [
    "text_pdf_en",
    "two_column_pdf",
    "sparse_canvas_pdf",
    "blank_pdf",
    "not_cv_pdf",
    "corrupt_pdf",
    "docx_cv",
    "vietnamese_docx_cv",
    "vietnamese_cv_txt",
    "mixed_language_cv_txt",
    "scanned_image_cv",
]
