"""Tests for the generalized filtered-export tools (Lane B).

``export_jobs`` / ``export_interviews`` / ``export_offers`` / ``export_events``:
registry + grant visibility, service-layer RBAC denial, cross-org isolation,
real .xlsx content sanity (localized headers, rows, filters, column subsets),
the frozen ``download`` render artifact, and the 24 h ChatExportFile expiry.

Offline/deterministic — no AI is involved in exports.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime, timedelta

import pytest
from app.modules.ai_assistant.application.native_loop import available_specs
from app.modules.ai_assistant.application.tools.dispatch import (
    SUPPORTED_TOOL_NAMES,
    dispatch_tool,
)
from app.modules.ai_assistant.application.tools.specs import TOOL_SPECS
from app.modules.ai_assistant.domain.models import ChatExportFile
from app.modules.documents.application import snapshot_service
from app.modules.opportunities.domain.event_models import Event, EventRegistration
from app.modules.recruitment.application import (
    access,
    apply_service,
    decision_service,
    interview_service,
    offer_service,
)

from tests.auth_utils import CTX
from tests.documents_utils import make_student
from tests.org_utils import add_member, make_org_with_admin
from tests.recruitment_utils import apply_payload, job_payload, make_builder_cv, publish_job

_EXPORT_TOOLS = ("export_jobs", "export_interviews", "export_offers", "export_events")


@pytest.fixture(autouse=True)
def _authorizer():
    access.install_authorizer()
    yield
    snapshot_service.set_snapshot_access_authorizer(None)


def _now() -> datetime:
    return datetime.now(tz=UTC)


async def _load_export(db, result: dict) -> ChatExportFile:
    """Resolve the stored file from the frozen download artifact."""

    render = result["render"]
    assert render["kind"] == "download"
    assert render["format"] == "xlsx"
    assert render["filename"] == result["filename"]
    export_id = uuid.UUID(render["download_path"].rsplit("/", 1)[-1])
    row = await db.get(ChatExportFile, export_id)
    assert row is not None
    return row


def _sheet(content: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(content)).active


def _header(ws) -> list[str]:
    return [c.value for c in ws[1] if c.value is not None]


async def _org_with_reviewed_app(db):
    """Partner org + published job + one reviewed application."""

    _pu, porg, partner = await make_org_with_admin(db, display_name="Exporter Co")
    _uu, _uorg, uni = await make_org_with_admin(db, org_type="university")
    job_id = await publish_job(db, partner_principal=partner, uni_principal=uni)
    _su, student = await make_student(db, prefix="exp_student")
    sel = await make_builder_cv(db, student=student)
    app = await apply_service.apply_to_job(
        db, principal=student, payload=apply_payload(job_id=job_id, cv_selection=sel), ctx=CTX
    )
    app_id = uuid.UUID(app["id"])
    await decision_service.review_application(db, principal=partner, application_id=app_id, ctx=CTX)
    return porg, partner, uni, job_id, app_id


# --------------------------------------------------------------------------- #
# Registry / grant visibility                                                  #
# --------------------------------------------------------------------------- #


def test_export_tools_registered_with_expected_grants() -> None:
    for name in _EXPORT_TOOLS:
        assert name in TOOL_SPECS
        assert name in SUPPORTED_TOOL_NAMES
        spec = TOOL_SPECS[name]
        assert spec.permission_class == "read_only"
        assert spec.persona == ["partner_user"]
        assert spec.audit_event_type.startswith("TOOL_")
    assert "jobs:export" in TOOL_SPECS["export_jobs"].required_permissions
    assert "applications:export" in TOOL_SPECS["export_interviews"].required_permissions
    assert "interviews:read" in TOOL_SPECS["export_interviews"].required_permissions
    assert "applications:export" in TOOL_SPECS["export_offers"].required_permissions
    assert "offers:create" in TOOL_SPECS["export_offers"].required_permissions
    assert "events:export" in TOOL_SPECS["export_events"].required_permissions


def test_catalog_gained_export_actions() -> None:
    from app.modules.organization.domain.catalog import is_catalog_permission

    assert is_catalog_permission("jobs", "export")
    assert is_catalog_permission("events", "export")
    assert is_catalog_permission("applications", "export")  # pre-existing


async def test_export_visibility_matches_grants(db_session) -> None:
    _u, porg, admin = await make_org_with_admin(db_session, display_name="Vis Co")
    # applications:export + interviews:read only.
    _mu, _m, interview_exporter = await add_member(
        db_session,
        org=porg,
        permissions=[("applications", "export"), ("interviews", "read")],
    )
    admin_tools = {s.name for s in available_specs(admin)}
    member_tools = {s.name for s in available_specs(interview_exporter)}

    assert set(_EXPORT_TOOLS) <= admin_tools
    assert "export_interviews" in member_tools
    assert "export_jobs" not in member_tools  # no jobs:export
    assert "export_offers" not in member_tools  # no offers:create
    assert "export_events" not in member_tools  # no events:export


# --------------------------------------------------------------------------- #
# export_jobs                                                                  #
# --------------------------------------------------------------------------- #


async def test_export_jobs_xlsx_content_and_artifact(db_session) -> None:
    from app.modules.opportunities.application import job_service

    _u, _org, admin = await make_org_with_admin(db_session, display_name="Jobs Co")
    await job_service.create_job(
        db_session, principal=admin, payload=job_payload("Backend Intern"), ctx=CTX
    )
    await job_service.create_job(
        db_session, principal=admin, payload=job_payload("Data Analyst"), ctx=CTX
    )

    res = await dispatch_tool("export_jobs", {}, session=db_session, principal=admin)
    assert res["ok"] is True
    assert res["row_count"] == 2

    export = await _load_export(db_session, res)
    assert export.mime.endswith("spreadsheetml.sheet")
    assert export.row_count == 2
    assert export.expires_at is not None
    expires = (
        export.expires_at.replace(tzinfo=UTC)
        if export.expires_at.tzinfo is None  # SQLite round-trips naive
        else export.expires_at
    )
    assert timedelta(hours=23) < (expires - _now()) <= timedelta(hours=24)

    ws = _sheet(export.content)
    assert _header(ws) == [
        "Tin tuyển dụng",
        "Trạng thái",
        "Số ứng viên",
        "Chưa xem",
        "Hạn ứng tuyển",
        "Ngày tạo",
        "Người phụ trách",
    ]
    titles = {ws.cell(row=r, column=1).value for r in (2, 3)}
    assert titles == {"Backend Intern", "Data Analyst"}
    # Status is the localized label, never the raw enum code.
    assert ws.cell(row=2, column=2).value not in ("draft", "pending_review")

    # Metadata-only analytics fact landed (kind + row_count, never row content).
    from app.modules.analytics.domain.models import AnalyticsEvent
    from sqlalchemy import select

    facts = (
        (
            await db_session.execute(
                select(AnalyticsEvent).where(
                    AnalyticsEvent.event_type == "ai.chat_export.generated"
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(facts) == 1
    assert facts[0].aggregate_id == export.id
    assert facts[0].properties["export"] == "jobs"
    assert facts[0].properties["row_count"] == 2
    assert "Backend Intern" not in str(facts[0].properties)


async def test_export_jobs_filters_and_column_subset(db_session) -> None:
    from app.modules.opportunities.application import job_service

    _u, _org, admin = await make_org_with_admin(db_session, display_name="Filter Co")
    await job_service.create_job(
        db_session, principal=admin, payload=job_payload("Only Draft"), ctx=CTX
    )

    # Status filter that matches nothing.
    res = await dispatch_tool(
        "export_jobs", {"status": "active"}, session=db_session, principal=admin
    )
    assert res["ok"] is True and res["row_count"] == 0

    # Date range excluding today.
    res = await dispatch_tool(
        "export_jobs",
        {"created_from": (_now() + timedelta(days=2)).strftime("%Y-%m-%d")},
        session=db_session,
        principal=admin,
    )
    assert res["ok"] is True and res["row_count"] == 0

    # Malformed date → deterministic user-safe error.
    res = await dispatch_tool(
        "export_jobs", {"created_from": "not-a-date"}, session=db_session, principal=admin
    )
    assert res == {"ok": False, "error": "invalid_date"}

    # Column subset preserved in order; unknown keys ignored.
    res = await dispatch_tool(
        "export_jobs",
        {"columns": ["owner", "title", "nope"]},
        session=db_session,
        principal=admin,
    )
    assert res["ok"] is True
    assert res["columns"] == ["owner", "title"]
    ws = _sheet((await _load_export(db_session, res)).content)
    assert _header(ws) == ["Người phụ trách", "Tin tuyển dụng"]


async def test_export_jobs_rbac_and_cross_org_isolation(db_session) -> None:
    from app.modules.opportunities.application import job_service

    _u, porg, admin = await make_org_with_admin(db_session, display_name="Org A")
    await job_service.create_job(
        db_session, principal=admin, payload=job_payload("A's Job"), ctx=CTX
    )

    # Member without jobs:export → service-layer denial (+ hidden from specs).
    _mu, _m, member = await add_member(db_session, org=porg, permissions=[("jobs", "read")])
    assert "export_jobs" not in {s.name for s in available_specs(member)}
    res = await dispatch_tool("export_jobs", {}, session=db_session, principal=member)
    assert res == {"ok": False, "error": "permission_denied"}

    # Another org's admin exports only their own (empty) list — never Org A's.
    _bu, _borg, admin_b = await make_org_with_admin(db_session, display_name="Org B")
    res_b = await dispatch_tool("export_jobs", {}, session=db_session, principal=admin_b)
    assert res_b["ok"] is True
    assert res_b["row_count"] == 0


# --------------------------------------------------------------------------- #
# export_interviews                                                            #
# --------------------------------------------------------------------------- #


async def test_export_interviews_content_no_meeting_link(db_session) -> None:
    porg, partner, _uni, _job_id, app_id = await _org_with_reviewed_app(db_session)
    _au, _am, assignee = await add_member(
        db_session,
        org=porg,
        permissions=[("applications", "read"), ("interviews", "read")],
    )
    secret_link = "https://meet.example/secret-room"
    await interview_service.schedule_interview(
        db_session,
        principal=partner,
        application_id=app_id,
        mode="online",
        scheduled_at=_now() + timedelta(days=2),
        assignee_ids=[assignee.user_id],
        meeting_link=secret_link,
        ctx=CTX,
    )

    res = await dispatch_tool(
        "export_interviews", {"scope": "upcoming"}, session=db_session, principal=partner
    )
    assert res["ok"] is True
    assert res["row_count"] == 1

    export = await _load_export(db_session, res)
    ws = _sheet(export.content)
    assert _header(ws) == [
        "Ứng viên",
        "Tin tuyển dụng",
        "Vòng",
        "Thời gian phỏng vấn",
        "Hình thức",
        "Người phỏng vấn",
        "Trạng thái",
    ]
    assert ws.cell(row=2, column=2).value == "Live Job"
    # The meeting link must never appear anywhere in the exported file.
    all_values = " ".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )
    assert secret_link not in all_values


async def test_export_interviews_filters(db_session) -> None:
    _porg, partner, _uni, job_id, app_id = await _org_with_reviewed_app(db_session)
    await interview_service.schedule_interview(
        db_session,
        principal=partner,
        application_id=app_id,
        mode="phone",
        scheduled_at=_now() + timedelta(days=3),
        assignee_ids=[partner.user_id],
        ctx=CTX,
    )

    # Past scope excludes the future interview.
    res = await dispatch_tool(
        "export_interviews", {"scope": "past"}, session=db_session, principal=partner
    )
    assert res["ok"] is True and res["row_count"] == 0

    # Date window around the interview includes it; a window before it excludes it.
    inside = await dispatch_tool(
        "export_interviews",
        {
            "scheduled_from": (_now() + timedelta(days=2)).strftime("%Y-%m-%d"),
            "scheduled_to": (_now() + timedelta(days=4)).strftime("%Y-%m-%d"),
        },
        session=db_session,
        principal=partner,
    )
    assert inside["ok"] is True and inside["row_count"] == 1
    before = await dispatch_tool(
        "export_interviews",
        {"scheduled_to": (_now() - timedelta(days=1)).strftime("%Y-%m-%d")},
        session=db_session,
        principal=partner,
    )
    assert before["ok"] is True and before["row_count"] == 0

    # Cross-org job_id → empty, never a tenant leak.
    res = await dispatch_tool(
        "export_interviews",
        {"job_id": str(uuid.uuid4())},
        session=db_session,
        principal=partner,
    )
    assert res["ok"] is True and res["row_count"] == 0


async def test_export_interviews_requires_board_read_grant(db_session) -> None:
    porg, _partner, _uni, _job_id, _app_id = await _org_with_reviewed_app(db_session)
    # applications:export WITHOUT interviews:read → the board's own gate denies.
    _mu, _m, member = await add_member(
        db_session, org=porg, permissions=[("applications", "export")]
    )
    res = await dispatch_tool("export_interviews", {}, session=db_session, principal=member)
    assert res == {"ok": False, "error": "permission_denied"}


# --------------------------------------------------------------------------- #
# export_offers                                                                #
# --------------------------------------------------------------------------- #


async def test_export_offers_content_salary_and_filters(db_session) -> None:
    _porg, partner, _uni, _job_id, app_id = await _org_with_reviewed_app(db_session)
    await offer_service.create_offer(
        db_session,
        principal=partner,
        application_id=app_id,
        position_title="Backend Engineer",
        expiry_date=_now() + timedelta(days=7),
        salary_amount=25_000_000,
        ctx=CTX,
    )

    res = await dispatch_tool("export_offers", {}, session=db_session, principal=partner)
    assert res["ok"] is True
    assert res["row_count"] == 1
    assert res["salary_included"] is True  # admin holds offers:create

    ws = _sheet((await _load_export(db_session, res)).content)
    assert _header(ws) == [
        "Ứng viên",
        "Tin tuyển dụng",
        "Vị trí",
        "Trạng thái",
        "Lương đề nghị",
        "Ngày gửi",
        "Ngày phản hồi",
    ]
    assert ws.cell(row=2, column=3).value == "Backend Engineer"
    assert "25,000,000" in str(ws.cell(row=2, column=5).value)
    # Status label, not the raw enum.
    assert ws.cell(row=2, column=4).value != "draft"

    # Status filter that matches nothing.
    res = await dispatch_tool(
        "export_offers", {"status": "accepted"}, session=db_session, principal=partner
    )
    assert res["ok"] is True and res["row_count"] == 0


async def test_export_offers_requires_offer_grant(db_session) -> None:
    porg, _partner, _uni, _job_id, _app_id = await _org_with_reviewed_app(db_session)
    # applications:export WITHOUT offers:create → the board's own gate denies
    # (same grant that unlocks comp on the offer board).
    _mu, _m, member = await add_member(
        db_session, org=porg, permissions=[("applications", "export")]
    )
    res = await dispatch_tool("export_offers", {}, session=db_session, principal=member)
    assert res == {"ok": False, "error": "permission_denied"}


async def test_export_offers_cross_org_isolation(db_session) -> None:
    _porg, partner, _uni, _job_id, app_id = await _org_with_reviewed_app(db_session)
    await offer_service.create_offer(
        db_session,
        principal=partner,
        application_id=app_id,
        position_title="Loot",
        expiry_date=_now() + timedelta(days=7),
        salary_amount=99_000_000,
        ctx=CTX,
    )
    _bu, _borg, admin_b = await make_org_with_admin(db_session, display_name="Rival Co")
    res = await dispatch_tool("export_offers", {}, session=db_session, principal=admin_b)
    assert res["ok"] is True
    assert res["row_count"] == 0


# --------------------------------------------------------------------------- #
# export_events                                                                #
# --------------------------------------------------------------------------- #


async def _seed_event_with_registrations(db, *, org_id, created_by) -> Event:
    event = Event(
        org_id=org_id,
        created_by=created_by,
        title="Career Fair 2026",
        slug=f"career-fair-{uuid.uuid4().hex[:8]}",
        description="Meet our teams.",
        event_type="career_fair",
        format="onsite",
        starts_at=_now() + timedelta(days=5),
        ends_at=_now() + timedelta(days=5, hours=4),
        capacity=100,
        status="published",
    )
    db.add(event)
    await db.flush()
    for st in ("confirmed", "attended", "waitlisted"):
        db.add(
            EventRegistration(event_id=event.id, user_id=created_by, status=st)
        )
    await db.commit()
    return event


async def test_export_events_aggregate_counts_no_pii(db_session) -> None:
    user, org, admin = await make_org_with_admin(db_session, display_name="Event Co")
    await _seed_event_with_registrations(db_session, org_id=org.id, created_by=user.id)

    res = await dispatch_tool("export_events", {}, session=db_session, principal=admin)
    assert res["ok"] is True
    assert res["row_count"] == 1

    ws = _sheet((await _load_export(db_session, res)).content)
    header = _header(ws)
    assert header[0] == "Sự kiện"
    assert "Đã đăng ký" in header and "Đã tham dự" in header
    row = {header[i]: ws.cell(row=2, column=i + 1).value for i in range(len(header))}
    assert row["Sự kiện"] == "Career Fair 2026"
    assert row["Đã đăng ký"] == 2  # confirmed + attended
    assert row["Danh sách chờ"] == 1
    assert row["Đã tham dự"] == 1
    # No attendee identity leaks into the file (emails, user ids).
    all_values = " ".join(
        str(c.value) for r in ws.iter_rows() for c in r if c.value is not None
    )
    assert "@" not in all_values
    assert str(user.id) not in all_values


async def test_export_events_rbac_and_isolation(db_session) -> None:
    user, org, _admin = await make_org_with_admin(db_session, display_name="Host Co")
    await _seed_event_with_registrations(db_session, org_id=org.id, created_by=user.id)

    # Member with events:read but not events:export → denied + hidden.
    _mu, _m, member = await add_member(db_session, org=org, permissions=[("events", "read")])
    assert "export_events" not in {s.name for s in available_specs(member)}
    res = await dispatch_tool("export_events", {}, session=db_session, principal=member)
    assert res == {"ok": False, "error": "permission_denied"}

    # Another org sees only its own (empty) event list.
    _bu, _borg, admin_b = await make_org_with_admin(db_session, display_name="Other Co")
    res_b = await dispatch_tool("export_events", {}, session=db_session, principal=admin_b)
    assert res_b["ok"] is True and res_b["row_count"] == 0


async def test_export_events_status_filter(db_session) -> None:
    user, org, admin = await make_org_with_admin(db_session, display_name="Status Co")
    await _seed_event_with_registrations(db_session, org_id=org.id, created_by=user.id)

    res = await dispatch_tool(
        "export_events", {"status": "draft"}, session=db_session, principal=admin
    )
    assert res["ok"] is True and res["row_count"] == 0
    res = await dispatch_tool(
        "export_events", {"status": "published"}, session=db_session, principal=admin
    )
    assert res["ok"] is True and res["row_count"] == 1
